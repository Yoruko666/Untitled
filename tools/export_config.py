"""Excel 配表导出工具。

把 Config/Table 下的 Excel 源表导出为 protobuf 二进制，
运行时通过 Addressables 加载 Assets/Configs/Database/Proto 下的 .bytes。

用法:
    python tools/export_config.py            # 导出所有表
    python tools/export_config.py Item       # 只导出指定表

表结构约定（每个 sheet 的前三行由本工具解析，不参与数据）:
    第 1 行   字段名，必须与 proto 字段一致
    第 2 行   类型说明，仅供策划阅读
    第 3 行   中文注释
    第 4 行起 数据，第一列为主键

同时会调用 protoc 刷新 Assets/Scripts/Data/Generated 下的 C# 代码。
"""

import importlib
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROTO_DIR = os.path.join(ROOT, "Assets", "Configs", "Database", "Proto")
TABLE_DIR = os.path.join(ROOT, "Config", "Table")
CSHARP_OUT_DIR = os.path.join(ROOT, "Assets", "Scripts", "Data", "Generated")
STUB_DIR = os.path.join(ROOT, "tools", ".generated")

HEADER_ROWS = 3

# 表名 -> 导出定义
TABLES = {
    "Item": {
        "workbook": "Item.xlsx",
        "sheet": "Item",
        "table": "ItemTable",
        "row": "Item",
        "repeated": "items",
        "key": "id",
        "output": "ItemTable.bytes",
        "module": "item_pb2",
    },
    "Bubble": {
        "workbook": "Bubble.xlsx",
        "sheet": "Bubble",
        "table": "BubbleTable",
        "row": "Bubble",
        "repeated": "bubbles",
        "key": "id",
        "output": "BubbleTable.bytes",
        "module": "bubble_pb2",
    },
    "Dialog": {
        "workbook": "Dialog.xlsx",
        "sheet": "Dialog",
        "table": "DialogTable",
        "row": "Dialog",
        "repeated": "dialogs",
        "key": "id",
        "output": "DialogTable.bytes",
        "module": "dialog_pb2",
    },
    "DialogOption": {
        "workbook": "Dialog.xlsx",
        "sheet": "DialogOption",
        "table": "DialogOptionTable",
        "row": "DialogOption",
        "repeated": "dialogOptions",
        "key": "id",
        "output": "DialogOptionTable.bytes",
        "module": "dialog_pb2",
    },
}


def run_protoc():
    protos = sorted(f for f in os.listdir(PROTO_DIR) if f.endswith(".proto"))
    if not protos:
        raise SystemExit(f"未找到 .proto 文件：{PROTO_DIR}")

    os.makedirs(STUB_DIR, exist_ok=True)
    os.makedirs(CSHARP_OUT_DIR, exist_ok=True)

    subprocess.run(
        [
            "protoc",
            f"-I{PROTO_DIR}",
            f"--python_out={STUB_DIR}",
            f"--csharp_out={CSHARP_OUT_DIR}",
        ]
        + protos,
        check=True,
    )
    normalize_newlines(CSHARP_OUT_DIR)
    sys.path.insert(0, STUB_DIR)


def normalize_newlines(directory):
    """protoc 在 Windows 上生成 CRLF，统一为 LF 以符合仓库规范。"""
    for folder, _, files in os.walk(directory):
        for name in files:
            if not name.endswith(".cs"):
                continue

            path = os.path.join(folder, name)
            with open(path, "r", encoding="utf-8", newline="") as handle:
                text = handle.read()
            if "\r\n" not in text:
                continue

            with open(path, "w", encoding="utf-8", newline="") as handle:
                handle.write(text.replace("\r\n", "\n"))


def read_records(path, sheet_name):
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"{path} 中不存在 sheet：{sheet_name}")

    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if len(rows) < HEADER_ROWS + 1:
        return [], []

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    missing = [i for i, name in enumerate(headers) if not name]
    if missing:
        raise SystemExit(f"{path} 的表头存在空列：{missing}")

    records = []
    for excel_row, raw in enumerate(rows[HEADER_ROWS:], start=HEADER_ROWS + 1):
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue

        record = {}
        for column, header in enumerate(headers):
            if column >= len(raw):
                continue
            value = raw[column]
            if value is None:
                continue
            if isinstance(value, str):
                value = value.strip()
                if not value:
                    continue
            record[header] = value
        records.append((excel_row, record))

    return headers, records


def normalize_enum_fields(row, record, output_name, excel_row):
    """允许策划写枚举短名（Ingredient），统一补成 proto 全名（ITEM_TYPE_INGREDIENT）。"""
    for field in row.DESCRIPTOR.fields:
        if field.type != field.TYPE_ENUM or field.name not in record:
            continue

        value = record[field.name]
        if not isinstance(value, str) or value.isdigit():
            continue

        names = [enum_value.name for enum_value in field.enum_type.values]
        if value in names:
            continue

        suffix = "_" + value.upper()
        matches = [name for name in names if name.endswith(suffix) or name.upper() == value.upper()]
        if len(matches) != 1:
            raise SystemExit(
                f"{output_name} 第 {excel_row} 行字段 {field.name} 的枚举值 {value} 无法唯一匹配，"
                f"请写全名，可选：{', '.join(names)}"
            )
        record[field.name] = matches[0]


def build_table(definition, headers, records, table_pb2, row_pb2):
    from google.protobuf import json_format

    key = definition["key"]
    if key not in headers:
        raise SystemExit(f"{definition['output']} 缺少主键列：{key}")

    table = table_pb2()
    repeated = getattr(table, definition["repeated"])
    seen = {}

    for excel_row, record in records:
        if key not in record:
            raise SystemExit(f"{definition['output']} 第 {excel_row} 行主键 {key} 为空")

        identity = record[key]
        if identity in seen:
            raise SystemExit(
                f"{definition['output']} 主键 {key}={identity} 重复："
                f"第 {seen[identity]} 行与第 {excel_row} 行"
            )
        seen[identity] = excel_row

        row = row_pb2()
        normalize_enum_fields(row, record, definition["output"], excel_row)
        try:
            json_format.ParseDict(record, row, ignore_unknown_fields=False)
        except Exception as error:  # noqa: BLE001 - 需要把行号带给策划
            raise SystemExit(f"{definition['output']} 第 {excel_row} 行解析失败：{error}") from error

        repeated.append(row)

    return table


def export(name, definition, stubs):
    workbook_path = os.path.join(TABLE_DIR, definition["workbook"])
    if not os.path.isfile(workbook_path):
        raise SystemExit(f"源表不存在：{workbook_path}")

    headers, records = read_records(workbook_path, definition["sheet"])
    table = build_table(
        definition,
        headers,
        records,
        getattr(stubs, definition["table"]),
        getattr(stubs, definition["row"]),
    )

    output_path = os.path.join(PROTO_DIR, definition["output"])
    with open(output_path, "wb") as handle:
        handle.write(table.SerializeToString())

    print(f"{name}: {len(records)} 行 -> {os.path.relpath(output_path, ROOT)}")


def main():
    requested = sys.argv[1:]
    unknown = [name for name in requested if name not in TABLES]
    if unknown:
        raise SystemExit(f"未定义的配表：{', '.join(unknown)}；可选：{', '.join(TABLES)}")

    run_protoc()

    modules = {}
    for name, definition in TABLES.items():
        if requested and name not in requested:
            continue

        module_name = definition["module"]
        if module_name not in modules:
            modules[module_name] = importlib.import_module(module_name)

        export(name, definition, modules[module_name])


if __name__ == "__main__":
    main()
